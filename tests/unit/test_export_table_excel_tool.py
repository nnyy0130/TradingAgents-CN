# -*- coding: utf-8 -*-
"""通用表格 Excel 导出工具测试。

背景：Excel 导出曾被绑死在估值工具（heavy 管线，2-6 分钟）上，
对话数据导出走估值管线导致 420 秒超时。本工具为确定性渲染：
Markdown 原文 → openpyxl，零 LLM 调用，秒级完成。
"""

import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.tools.implementations.assistant_ops.export_table_excel_tool import (
    export_table_excel_tool,
    _parse_cell,
    _parse_markdown_tables,
    _sanitize_sheet_name,
)

_TABLE = (
    "## 月度销量预测\n"
    "| 月份 | 问界M9 | 问界M7 |\n"
    "|---|---:|---:|\n"
    "| 2026-09 | 10,500 | 8,000 |\n"
    "| 2026-10 | 11,000 | 8,200 |\n"
)


# ---------- 纯函数 ----------

def test_parse_cell_types():
    assert _parse_cell("10,500") == (10500.0, "#,##0.##")
    assert _parse_cell("-3") == (-3, "#,##0")
    assert _parse_cell("12.5%") == (pytest.approx(0.125), "0.00%")
    assert _parse_cell("3.14") == (3.14, "0.##")
    assert _parse_cell("问界M9")[0] == "问界M9"


def test_parse_markdown_tables_multi_sheet():
    md = (
        "## 表A\n| a | b |\n|---|---|\n| 1 | 2 |\n\n"
        "## 表B\n| x |\n|---|\n| 9 |\n"
    )
    parsed = _parse_markdown_tables(md)
    assert len(parsed) == 2
    assert parsed[0][0] == "表A"
    assert parsed[0][1] == [["a", "b"], ["1", "2"]]
    assert parsed[1][0] == "表B"


def test_parse_markdown_tables_no_table():
    assert _parse_markdown_tables("没有表格的普通文本") == []


def test_sanitize_sheet_name_dedup():
    used = set()
    n1 = _sanitize_sheet_name("销量表", used)
    n2 = _sanitize_sheet_name("销量表", used)
    assert n1 != n2
    assert _sanitize_sheet_name("a/b\\c?*", set()) == "a_b_c__"


# ---------- 工具集成（真实渲染文件） ----------

def test_export_creates_excel_file(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    result = export_table_excel_tool.invoke({
        "title": "赛力斯营收预测",
        "markdown_tables": _TABLE,
        "notes": "数据来源：懂车帝销量榜，ASP 取指导价中位数",
    })

    assert "已导出" in result
    assert "/api/assistant/deliverables/table_" in result
    assert "1 个 Sheet" in result
    assert "不构成投资建议" in result

    m = re.search(r"deliverables/(table_[^)?]+\.xlsx)", result)
    assert m, f"下载链接缺失: {result}"
    # 链接中文件名经 URL 编码，需解码后核对落盘文件
    from urllib.parse import unquote

    f = Path("data/deliverables") / unquote(m.group(1))
    assert f.exists()
    # 签名下载链接：浏览器直接点击（无 JWT 头）可下载
    assert "exp=" in result and "sig=" in result

    wb = load_workbook(str(f))
    ws = wb["月度销量预测"]
    assert ws.cell(1, 1).value == "月份"
    assert ws.cell(2, 2).value == 10500.0
    assert ws.cell(3, 3).value == 8200


def test_export_numbers_verbatim(tmp_path, monkeypatch):
    """数字逐字透传：Excel 中的值与 Markdown 原文一致（不改写）。"""
    monkeypatch.chdir(tmp_path)
    md = (
        "| 指标 | 2026E |\n|---|---:|\n"
        "| Q2 校准误差 | 0.06% |\n"
        "| 测算营收（亿） | 317.29 |\n"
    )
    result = export_table_excel_tool.invoke({
        "title": "模型校准", "markdown_tables": md,
    })
    assert "已导出" in result

    m = re.search(r"deliverables/(table_[^)?]+\.xlsx)", result)
    from urllib.parse import unquote

    wb = load_workbook(str(Path("data/deliverables") / unquote(m.group(1))))
    ws = wb.active
    assert ws.cell(2, 2).value == pytest.approx(0.0006)
    assert ws.cell(3, 2).value == 317.29


def test_export_no_table_returns_failure(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    result = export_table_excel_tool.invoke({
        "title": "空", "markdown_tables": "这里没有表格",
    })
    assert "导出失败" in result
    assert "未在传入内容中解析到" in result


def test_tool_registered_in_registry():
    from core.tools import get_tool_registry
    meta = get_tool_registry().get("export_table_excel_tool")
    assert meta is not None, "export_table_excel_tool 必须自动注册到工具注册表"
    # light 层级（15s）：纯 CPU 渲染，与估值工具的 heavy（180s）形成对比
    assert meta.timeout_tier == "light"
