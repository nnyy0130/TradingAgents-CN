# -*- coding: utf-8 -*-
"""估值测算 Excel 交付物单测（A10-D1）

覆盖：合法规格渲染+校验、公式格存在性、规格阻断规则（PE 负盈利/缺字段/零倍数）、
文件篡改检测（死值替换公式 / 输入值漂移）、技能入口往返。
"""

import json

import pytest
from openpyxl import load_workbook

from core.deliverables.valuation_excel import (
    ValuationSpecError,
    generate_valuation_excel,
    render_valuation_excel,
    validate_spec,
    validate_valuation_excel,
)

# ---- 测试规格（数值为演示口径，非投资结论）----

VALID_SPEC = {
    "ticker": "601127",
    "stock_name": "赛力斯",
    "analysis_date": "2026-09-10",
    "share_count": 15.1,
    "current_price": 130.0,
    "earnings_forecast": {
        "history": [
            {"period": "2024A", "revenue": 1451.8, "net_profit": 59.5},
            {"period": "2025A", "revenue": 1900.0, "net_profit": 85.0},
        ],
        "assumptions": [
            {"period": "2026E", "revenue_growth": 0.30, "net_margin": 0.065,
             "rationale": "新车型放量与产能爬坡"},
            {"period": "2027E", "revenue_growth": 0.20, "net_margin": 0.07,
             "rationale": "规模效应下净利率继续改善"},
        ],
    },
    "valuation": {
        "scenarios": [
            {"name": "悲观", "method": "PE", "multiple": 12, "eps_period": "2026E",
             "benchmark_basis": "同业低端 12 倍",
             "applicability_check": {"passed": True, "notes": ["归母净利润为正"]}},
            {"name": "中性", "method": "PE", "multiple": 22, "eps_period": "2026E",
             "benchmark_basis": "同业 22-28 倍区间下沿",
             "applicability_check": {"passed": True, "notes": ["归母净利润为正"]}},
            {"name": "乐观", "method": "PE", "multiple": 28, "eps_period": "2027E",
             "benchmark_basis": "同业 22-28 倍区间上沿",
             "applicability_check": {"passed": True, "notes": ["归母净利润为正"]}},
        ],
        "risk_notes": ["测算为情景假设，不构成投资建议"],
    },
}


def test_valid_spec_passes_spec_validation():
    assert validate_spec(VALID_SPEC) == []


def test_render_and_validate_roundtrip(tmp_path):
    path = str(tmp_path / "valuation.xlsx")
    out = render_valuation_excel(VALID_SPEC, path)

    import os
    assert os.path.exists(out)

    report = validate_valuation_excel(VALID_SPEC, out)
    assert report["passed"], report["errors"]
    assert report["errors"] == []

    # Python 重算期望值（计算自洽的数值锚点）
    # 2026E: 营收 1900*1.3=2470, 净利 2470*0.065=160.55, EPS=160.55/15.1
    eps_2026 = 2470.0 * 0.065 / 15.1
    # 2027E: 营收 2470*1.2=2964, 净利 2964*0.07=207.48, EPS=207.48/15.1
    eps_2027 = 2964.0 * 0.07 / 15.1
    summary = report["summary"]["scenarios"]
    assert summary["悲观"]["eps"] == pytest.approx(eps_2026, abs=1e-3)
    assert summary["中性"]["target_price"] == pytest.approx(eps_2026 * 22, abs=0.1)
    assert summary["乐观"]["target_price"] == pytest.approx(eps_2027 * 28, abs=0.1)


def test_forecast_cells_are_formulas(tmp_path):
    path = str(tmp_path / "valuation.xlsx")
    render_valuation_excel(VALID_SPEC, path)

    wb = load_workbook(path)
    ws = wb["盈利推演"]
    # 布局：行4=2024A，行5=2025A，行6=2026E，行7=2027E
    assert ws["B6"].value == "=B5*(1+C6)"
    assert ws["E6"].value == "=B6*D6"
    assert ws["B7"].value == "=B6*(1+C7)"
    # 假设格是输入值（非公式）
    assert ws["C6"].value == 0.30
    assert ws["D7"].value == 0.07

    ws_v = wb["估值情景"]
    # 行6=悲观：EPS 引用盈利推演行6（2026E），目标价 = EPS × 倍数
    assert ws_v["E6"].value == "=盈利推演!E6/$B$2"
    assert ws_v["F6"].value == "=E6*C6"
    assert ws_v["B2"].value == 15.1


def test_pe_with_negative_profit_blocked():
    spec = json.loads(json.dumps(VALID_SPEC))
    spec["earnings_forecast"]["history"].insert(
        0, {"period": "2023A", "revenue": 358.4, "net_profit": -24.5}
    )
    spec["valuation"]["scenarios"][0]["eps_period"] = "2023A"

    errors = validate_spec(spec)
    assert any("PE" in e and "正" in e for e in errors), errors

    with pytest.raises(ValuationSpecError, match="PE"):
        render_valuation_excel(spec, "unused.xlsx")


def test_missing_required_field_blocked():
    spec = json.loads(json.dumps(VALID_SPEC))
    del spec["valuation"]["scenarios"][0]["benchmark_basis"]

    errors = validate_spec(spec)
    assert any("benchmark_basis" in e for e in errors), errors


def test_zero_multiple_and_margin_blocked():
    spec = json.loads(json.dumps(VALID_SPEC))
    spec["valuation"]["scenarios"][0]["multiple"] = 0
    assert any("倍数为零" in e for e in validate_spec(spec))

    spec = json.loads(json.dumps(VALID_SPEC))
    spec["earnings_forecast"]["assumptions"][0]["net_margin"] = 0
    assert any("净利率为零" in e for e in validate_spec(spec))


def test_tampered_dead_value_detected(tmp_path):
    """预测行公式被替换为死值 → 校验拦截"""
    path = str(tmp_path / "valuation.xlsx")
    render_valuation_excel(VALID_SPEC, path)

    wb = load_workbook(path)
    wb["盈利推演"]["B6"] = 2470.0  # 死值替换公式
    tampered = str(tmp_path / "tampered_dead.xlsx")
    wb.save(tampered)

    report = validate_valuation_excel(VALID_SPEC, tampered)
    assert not report["passed"]
    assert any("不是公式格" in e for e in report["errors"]), report["errors"]


def test_tampered_input_value_detected(tmp_path):
    """假设输入格被改 → 与规格不一致，校验拦截"""
    path = str(tmp_path / "valuation.xlsx")
    render_valuation_excel(VALID_SPEC, path)

    wb = load_workbook(path)
    wb["盈利推演"]["C6"] = 0.50  # 增速假设被改
    tampered = str(tmp_path / "tampered_input.xlsx")
    wb.save(tampered)

    report = validate_valuation_excel(VALID_SPEC, tampered)
    assert not report["passed"]
    assert any("增速输入值与规格不一致" in e for e in report["errors"]), report["errors"]


def test_skill_entry_roundtrip(tmp_path):
    result_json = generate_valuation_excel(
        json.dumps(VALID_SPEC, ensure_ascii=False),
        output_path=str(tmp_path / "skill_out.xlsx"),
    )
    result = json.loads(result_json)
    assert result["success"], result
    assert result["path"].endswith(".xlsx")
    assert result["summary"]["scenarios"]["中性"]["eps"] > 0


def test_skill_entry_invalid_spec():
    bad = json.loads(json.dumps(VALID_SPEC))
    bad["valuation"]["scenarios"][0]["multiple"] = -5
    result = json.loads(generate_valuation_excel(json.dumps(bad)))
    assert not result["success"]
    assert result["stage"] == "spec_validation"
